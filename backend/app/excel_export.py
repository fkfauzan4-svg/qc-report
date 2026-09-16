"""Builds the downloadable Excel workbook (Summary, BOQ & Quotation, Detailed
BOM Breakdown, Hardware & Accessories) from the extracted/edited payload.

Mirrors the structure of the reference workbook this app was modeled after,
using the same deep-slate/amber styling as the company's existing QC report
template so exports look consistent."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schema import BOM_COLUMNS, HARDWARE_COLUMNS, ITEM_COLUMNS

COLOR_HEADER_FILL = "2C3E50"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_SECTION_FILL = "EAEDED"
COLOR_TOTAL_FILL = "FCF3CF"
COLOR_BORDER = "BDC3C7"
COLOR_INPUT_TEXT = "0000FF"

FONT_NAME = "Segoe UI"

font_title = Font(name=FONT_NAME, size=16, bold=True, color=COLOR_HEADER_FILL)
font_section = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_HEADER_FILL)
font_header = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_HEADER_TEXT)
font_body = Font(name=FONT_NAME, size=10)
font_body_bold = Font(name=FONT_NAME, size=10, bold=True)
font_input = Font(name=FONT_NAME, size=10, color=COLOR_INPUT_TEXT)
font_total = Font(name=FONT_NAME, size=11, bold=True)

fill_header = PatternFill(start_color=COLOR_HEADER_FILL, end_color=COLOR_HEADER_FILL, fill_type="solid")
fill_section = PatternFill(start_color=COLOR_SECTION_FILL, end_color=COLOR_SECTION_FILL, fill_type="solid")
fill_total = PatternFill(start_color=COLOR_TOTAL_FILL, end_color=COLOR_TOTAL_FILL, fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER),
)


def _num(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _autosize(ws, min_width=10, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or isinstance(cell.value, str) and cell.value.startswith("="):
                continue
            col = cell.column_letter
            length = len(str(cell.value))
            widths[col] = max(widths.get(col, 0), length)
    for col, length in widths.items():
        ws.column_dimensions[col].width = max(min_width, min(max_width, length + 3))


def build_workbook(payload: dict) -> bytes:
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    vat_ref = _build_summary_sheet(ws_summary, payload)

    ws_boq = wb.create_sheet("BOQ & Quotation")
    _build_boq_sheet(ws_boq, payload, vat_ref)

    ws_bom = wb.create_sheet("Detailed BOM Breakdown")
    _build_bom_sheet(ws_bom, payload)

    ws_hw = wb.create_sheet("Hardware & Accessories")
    _build_hardware_sheet(ws_hw, payload)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_summary_sheet(ws, payload):
    project = payload.get("project", {})

    ws["A1"] = "BOM & COST ESTIMATION REPORT"
    ws["A1"].font = font_title
    ws.row_dimensions[1].height = 26

    rows = [
        ("Project Name:", project.get("projectName", ""), "Document Ref:", project.get("documentRef", "")),
        ("Client:", project.get("client", ""), "Manufacturer:", project.get("manufacturer", "")),
        ("Date:", project.get("date", ""), "Currency:", project.get("currency", "SAR")),
    ]
    r = 3
    for label1, val1, label2, val2 in rows:
        ws.cell(row=r, column=1, value=label1).font = font_body_bold
        ws.cell(row=r, column=2, value=val1).font = font_body
        ws.cell(row=r, column=4, value=label2).font = font_body_bold
        ws.cell(row=r, column=5, value=val2).font = font_body
        r += 1

    ws.cell(row=r, column=1, value="VAT Rate:").font = font_body_bold
    vat_cell = ws.cell(row=r, column=2, value=_num(project.get("vatRate", 0.15), 0.15))
    vat_cell.font = font_input
    vat_cell.number_format = "0.0%"
    vat_ref = f"'{ws.title}'!$B${r}"

    r += 2
    ws.cell(row=r, column=1, value="Notes:").font = font_body_bold
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(row=r, column=1, value=project.get("notes", "")).font = font_body
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 40

    r += 2
    ws.cell(row=r, column=1, value="ASSUMPTIONS & NOTES FROM AUTOMATED DRAWING ANALYSIS").font = font_section
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(row=r, column=1).fill = fill_section
    r += 1
    assumptions = payload.get("assumptions") or []
    if not assumptions:
        assumptions = ["No assumptions were recorded for this analysis."]
    for note in assumptions:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        cell = ws.cell(row=r, column=1, value=f"• {note}")
        cell.font = font_body
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 18
        r += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 26

    return vat_ref


def _build_boq_sheet(ws, payload, vat_ref):
    ws["A1"] = "COMMERCIAL BOQ & QUOTATION"
    ws["A1"].font = font_title
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ITEM_COLUMNS) + 1)

    header_row = 3
    headers = [label for _, label in ITEM_COLUMNS] + ["Total Cost"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 26

    items = payload.get("items") or []
    qty_col = [i for i, (key, _) in enumerate(ITEM_COLUMNS) if key == "qty"][0] + 1
    cost_col = [i for i, (key, _) in enumerate(ITEM_COLUMNS) if key == "unitCost"][0] + 1
    total_col = len(ITEM_COLUMNS) + 1

    r = header_row + 1
    first_data_row = r
    for item in items:
        for idx, (key, _) in enumerate(ITEM_COLUMNS, start=1):
            value = item.get(key, "")
            cell = ws.cell(row=r, column=idx, value=value)
            cell.border = thin_border
            if key in ("qty", "unitCost"):
                cell.font = font_input
                cell.number_format = "#,##0.00"
            else:
                cell.font = font_body
        qty_letter = get_column_letter(qty_col)
        cost_letter = get_column_letter(cost_col)
        total_cell = ws.cell(row=r, column=total_col, value=f"={qty_letter}{r}*{cost_letter}{r}")
        total_cell.font = font_body_bold
        total_cell.number_format = "#,##0.00"
        total_cell.border = thin_border
        r += 1
    last_data_row = r - 1
    if last_data_row < first_data_row:
        last_data_row = first_data_row  # empty table, formulas still valid

    total_letter = get_column_letter(total_col)

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col - 1)
    ws.cell(row=r, column=1, value="SUBTOTAL (Excl. VAT):").font = font_total
    subtotal_cell = ws.cell(row=r, column=total_col, value=f"=SUM({total_letter}{first_data_row}:{total_letter}{last_data_row})")
    subtotal_cell.font = font_total
    subtotal_cell.number_format = "#,##0.00"
    subtotal_cell.fill = fill_total
    subtotal_row = r

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col - 1)
    ws.cell(row=r, column=1, value="VAT:").font = font_total
    vat_cell = ws.cell(row=r, column=total_col, value=f"={total_letter}{subtotal_row}*{vat_ref}")
    vat_cell.font = font_total
    vat_cell.number_format = "#,##0.00"
    vat_cell.fill = fill_total
    vat_row = r

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col - 1)
    ws.cell(row=r, column=1, value="GRAND TOTAL (Incl. VAT):").font = font_total
    grand_cell = ws.cell(row=r, column=total_col, value=f"={total_letter}{subtotal_row}+{total_letter}{vat_row}")
    grand_cell.font = font_total
    grand_cell.number_format = "#,##0.00"
    grand_cell.fill = fill_total

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _autosize(ws, min_width=12, max_width=55)
    ws.column_dimensions[get_column_letter(3)].width = 55


def _build_bom_sheet(ws, payload):
    ws["A1"] = "DETAILED BOM BREAKDOWN"
    ws["A1"].font = font_title
    ws.row_dimensions[1].height = 26

    header_row = 3
    for idx, (_, label) in enumerate(BOM_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=label)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 26

    rows = payload.get("bom") or []
    r = header_row + 1
    for row_data in rows:
        for idx, (key, _) in enumerate(BOM_COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx, value=row_data.get(key, ""))
            cell.border = thin_border
            cell.font = font_body
            if key in ("qtyPerUnit", "totalQty"):
                cell.number_format = "#,##0.00"
        r += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _autosize(ws, min_width=12, max_width=45)


def _build_hardware_sheet(ws, payload):
    ws["A1"] = "HARDWARE & ACCESSORIES SCHEDULE"
    ws["A1"].font = font_title
    ws.row_dimensions[1].height = 26

    note_row = 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(HARDWARE_COLUMNS) + 1)
    note_cell = ws.cell(
        row=note_row,
        column=1,
        value=(
            "Reference breakdown for production planning. Costs here are already "
            "included within each item's unit price on the BOQ sheet — do not add "
            "this total on top of the BOQ grand total."
        ),
    )
    note_cell.font = Font(name=FONT_NAME, size=9, italic=True, color="7B7D7D")

    header_row = 4
    headers = [label for _, label in HARDWARE_COLUMNS] + ["Total Cost"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 26

    rows = payload.get("hardware") or []
    qty_col = [i for i, (key, _) in enumerate(HARDWARE_COLUMNS) if key == "totalQty"][0] + 1
    cost_col = [i for i, (key, _) in enumerate(HARDWARE_COLUMNS) if key == "unitCost"][0] + 1
    total_col = len(HARDWARE_COLUMNS) + 1

    r = header_row + 1
    first_data_row = r
    for row_data in rows:
        for idx, (key, _) in enumerate(HARDWARE_COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx, value=row_data.get(key, ""))
            cell.border = thin_border
            if key in ("qtyPerUnit", "totalQty", "unitCost"):
                cell.font = font_input
                cell.number_format = "#,##0.00"
            else:
                cell.font = font_body
        qty_letter = get_column_letter(qty_col)
        cost_letter = get_column_letter(cost_col)
        total_cell = ws.cell(row=r, column=total_col, value=f"={qty_letter}{r}*{cost_letter}{r}")
        total_cell.font = font_body_bold
        total_cell.number_format = "#,##0.00"
        total_cell.border = thin_border
        r += 1
    last_data_row = r - 1
    if last_data_row < first_data_row:
        last_data_row = first_data_row

    total_letter = get_column_letter(total_col)
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col - 1)
    ws.cell(row=r, column=1, value="TOTAL HARDWARE & ACCESSORIES COST:").font = font_total
    total_cell = ws.cell(row=r, column=total_col, value=f"=SUM({total_letter}{first_data_row}:{total_letter}{last_data_row})")
    total_cell.font = font_total
    total_cell.number_format = "#,##0.00"
    total_cell.fill = fill_total

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    _autosize(ws, min_width=12, max_width=45)
